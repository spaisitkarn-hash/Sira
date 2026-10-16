"""Build the DMS Introduction & 5 Components training deck.

    python3 build_deck.py [output.pptx]

Requires: python-pptx, cairosvg, pillow, lxml. Icons: Material Symbols Rounded SVGs in assets/icons
(Apache-2.0). Font: Bai Jamjuree (OFL) in fonts/ — install it on the presenting PC, or embed it when saving.
"""
import os
import sys

from pptx import Presentation

import dms_lib as L
from slides_components import SLIDES as COMPONENT_SLIDES
from slides_intro import SLIDES as INTRO_SLIDES


def build(out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = L.In(L.W), L.In(L.H)
    blank = prs.slide_layouts[6]
    slides = INTRO_SLIDES + COMPONENT_SLIDES
    total = len(slides)
    for n, fn in enumerate(slides, 1):
        sl = prs.slides.add_slide(blank)
        fn(sl, n, total)
        dark = getattr(fn, "dark", False)
        if n > 1:
            L.footer(sl, n, total, accent=getattr(fn, "accent", L.NAVY), dark=dark)
    prs.core_properties.title = "Daily Management System — DMS Introduction & 5 Components"
    prs.core_properties.author = "Sappe Continuous Improvement"
    prs.core_properties.subject = "DMS Training"
    prs.save(out)
    return total


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "DMS_Introduction_5_Components.pptx")
    print(f"{build(out)} slides -> {out}")
