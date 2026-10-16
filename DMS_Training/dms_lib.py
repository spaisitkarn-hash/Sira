"""Design system + drawing helpers for the DMS training deck (python-pptx).

All geometry is in inches on a 13.333 x 7.5 in (16:9) canvas.
"""
import math
import os
import re
import tempfile

import cairosvg
from lxml import etree
from PIL import ImageFont
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ICON_DIR = os.environ.get("DMS_ICON_DIR", os.path.join(HERE, "assets", "icons"))
FONT_DIR = os.path.join(HERE, "fonts")
_PNG_CACHE = tempfile.mkdtemp(prefix="dms_icons_")

W, H = 13.333, 7.5
ML, MR = 0.65, 12.68          # content left / right edge
CW = MR - ML                  # content width

# ---------------------------------------------------------------- palette
NAVY = "14275A"
NAVY_D = "0E1D47"
INK = "1E2B4A"
SLATE = "55627A"
MUTED = "8A94A8"
LINE = "DDE3EE"
BG = "F3F6FB"
WHITE = "FFFFFF"
MINT = "7FDBC0"
AMBER = "E8A317"

GREEN, GREEN_D, GREEN_L = "0E8C6C", "0B5A45", "E3F3EE"
BLUE, BLUE_D, BLUE_L = "1A66D9", "0D3E8F", "E6EFFC"
ORANGE, ORANGE_D, ORANGE_L = "D9730D", "9A4A05", "FDF0E1"
RED, RED_D, RED_L = "C8361F", "8E2414", "FBE7E3"
INDIGO, INDIGO_D, INDIGO_L = "3F51B5", "27347E", "E9EBF7"

COMP = {
    1: dict(c=GREEN, d=GREEN_D, l=GREEN_L, icon="account_tree", name="KPI Cascade"),
    2: dict(c=BLUE, d=BLUE_D, l=BLUE_L, icon="groups", name="Purposeful Meeting Structure"),
    3: dict(c=ORANGE, d=ORANGE_D, l=ORANGE_L, icon="notification_important", name="Deviation & Escalation"),
    4: dict(c=RED, d=RED_D, l=RED_L, icon="troubleshoot", name="RCA Management"),
    5: dict(c=INDIGO, d=INDIGO_D, l=INDIGO_L, icon="fact_check", name="Process Confirmation"),
}

F = "Bai Jamjuree"
FM = "Bai Jamjuree Medium"
FSB = "Bai Jamjuree SemiBold"
_FONT_FILES = {F: "BaiJamjuree-Regular.ttf", FM: "BaiJamjuree-Medium.ttf",
               FSB: "BaiJamjuree-SemiBold.ttf", "bold": "BaiJamjuree-Bold.ttf"}


def In(v):
    return Emu(int(round(v * 914400)))


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- measuring
_font_cache = {}


def text_w(text, size, font=F, bold=False):
    """Approximate rendered width of a single line, in inches."""
    key = "bold" if bold and font == F else font
    if key not in _font_cache:
        _font_cache[key] = ImageFont.truetype(os.path.join(FONT_DIR, _FONT_FILES[key]), 200)
    return _font_cache[key].getlength(text) * size / 200 / 72


# ---------------------------------------------------------------- xml helpers
def _set_alpha(fill_parent_el, alpha):
    clr = fill_parent_el.find(".//" + qn("a:srgbClr"))
    if clr is not None:
        for a in clr.findall(qn("a:alpha")):
            clr.remove(a)
        etree.SubElement(clr, qn("a:alpha"), val=str(int(alpha * 100000)))


def no_shadow(shp):
    shp.shadow.inherit = False


def shadow(shp, blur=14, dist=3, alpha=0.10, color=NAVY):
    spPr = shp._element.spPr
    eff = spPr.find(qn("a:effectLst"))
    if eff is None:
        eff = etree.SubElement(spPr, qn("a:effectLst"))
    for c in list(eff):
        eff.remove(c)
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(int(Pt(blur))), dist=str(int(Pt(dist))),
                          dir="5400000", algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val=color)
    etree.SubElement(clr, qn("a:alpha"), val=str(int(alpha * 100000)))


def _set_adj(shp, **vals):
    """Set raw preset-geometry guide values (e.g. adj1=10800000)."""
    prst = shp._element.spPr.find(qn("a:prstGeom"))
    av = prst.find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(prst, qn("a:avLst"))
    for c in list(av):
        av.remove(c)
    for k, v in vals.items():
        etree.SubElement(av, qn("a:gd"), name=k, fmla=f"val {int(v)}")


# ---------------------------------------------------------------- text
def _apply_run(run, size, color, bold, font, spc=None, alpha=None):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        for e in rPr.findall(qn(tag)):
            rPr.remove(e)
    for tag in ("a:latin", "a:ea", "a:cs"):
        etree.SubElement(rPr, qn(tag), typeface=font)
    if spc is not None:
        rPr.set("spc", str(int(spc * 100)))
    if alpha is not None:
        _set_alpha(rPr, alpha)


def set_text(tf, content, size=14, color=INK, bold=False, font=F, align="l", anchor="t",
             ls=None, after=0, spc=None, alpha=None, margin=0.0, wrap=True):
    """content: str (\\n = new paragraph) or list of paragraphs; a paragraph is a str
    or a list of runs; a run is a str or (text, {size,color,bold,font}) tuple.
    A paragraph may also be a dict {"runs": [...], "align":..., "after":..., "size":...}."""
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    m = In(margin)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    paras = content.split("\n") if isinstance(content, str) else content
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pal, paf, psz = align, after, size
        if isinstance(p, dict):
            pal, paf, psz = p.get("align", align), p.get("after", after), p.get("size", size)
            p = p["runs"]
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[pal]
        if ls:
            para.line_spacing = ls
        para.space_after = Pt(paf)
        runs = [p] if isinstance(p, (str, tuple)) else p
        for r in runs:
            t, o = (r, {}) if isinstance(r, str) else r
            run = para.add_run()
            run.text = t
            _apply_run(run, o.get("size", psz), o.get("color", color), o.get("bold", bold),
                       o.get("font", font), o.get("spc", spc), o.get("alpha", alpha))
    return tf


def txt(sl, x, y, w, h, content, **kw):
    tb = sl.shapes.add_textbox(In(x), In(y), In(w), In(h))
    set_text(tb.text_frame, content, **kw)
    return tb


# ---------------------------------------------------------------- shapes
def rect(sl, x, y, w, h, fill=None, line=None, lw=0.75, r=None, grad=None, angle=90,
         alpha=None, sh=False, shape=None, dash=False, text=None, **tkw):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE)
    s = sl.shapes.add_shape(kind, In(x), In(y), In(w), In(h))
    if r and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, r / min(w, h))
    if grad:
        s.fill.gradient()
        s.fill.gradient_angle = angle
        st = s.fill.gradient_stops
        st[0].color.rgb, st[0].position = rgb(grad[0]), 0
        st[1].color.rgb, st[1].position = rgb(grad[1]), 1.0
    elif fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            _set_alpha(s._element.spPr, alpha)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        s.line.fill.background()
    no_shadow(s)
    if sh:
        shadow(s)
    if text is not None:
        tkw.setdefault("margin", 0.08)
        set_text(s.text_frame, text, **tkw)
    else:
        s.text_frame.text = ""
    return s


def circle(sl, cx, cy, d, fill=None, line=None, lw=1, **kw):
    return rect(sl, cx - d / 2, cy - d / 2, d, d, fill=fill, line=line, lw=lw, shape=MSO_SHAPE.OVAL, **kw)


def line(sl, x1, y1, x2, y2, color=MUTED, w=1.5, arrow=False, head=False, dash=False, alpha=None):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, In(x1), In(y1), In(x2), In(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    ln = c.line._get_or_add_ln()
    if alpha is not None:
        _set_alpha(ln, alpha)
    if head:
        etree.SubElement(ln, qn("a:headEnd"), type="triangle", w="med", len="med")
    if arrow:
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


def poly(sl, pts, color=MUTED, w=1.5, arrow=True, dash=False):
    """Orthogonal/poly connector built from straight segments; arrow on the last one."""
    for i in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        line(sl, x1, y1, x2, y2, color, w, arrow=arrow and i == len(pts) - 2, dash=dash)


def arc(sl, cx, cy, r, start, end, color, w_pt):
    """Outline arc (angles in screen degrees, 0 = 3 o'clock, clockwise)."""
    s = sl.shapes.add_shape(MSO_SHAPE.ARC, In(cx - r), In(cy - r), In(2 * r), In(2 * r))
    _set_adj(s, adj1=start * 60000, adj2=end * 60000)
    s.fill.background()
    s.line.color.rgb = rgb(color)
    s.line.width = Pt(w_pt)
    no_shadow(s)
    return s


# ---------------------------------------------------------------- icons
def _icon_png(name, color):
    out = os.path.join(_PNG_CACHE, f"{name}_{color}.png")
    if os.path.exists(out):
        return out
    path = os.path.join(ICON_DIR, f"{name}-fill.svg")
    if not os.path.exists(path):
        path = os.path.join(ICON_DIR, f"{name}.svg")
    with open(path, encoding="utf-8") as fh:
        svg = fh.read()
    svg = re.sub(r"<svg ", f'<svg fill="#{color}" ', svg, count=1)
    cairosvg.svg2png(bytestring=svg.encode(), write_to=out, output_width=256, output_height=256)
    return out


def icon(sl, name, x, y, s, color=NAVY):
    return sl.shapes.add_picture(_icon_png(name, color), In(x), In(y), In(s), In(s))


def icon_c(sl, name, cx, cy, s, color=NAVY):
    return icon(sl, name, cx - s / 2, cy - s / 2, s, color)


def icon_badge(sl, name, x, y, d, bg, fg=WHITE, round_=True, r=0.12, ratio=0.55, sh=False):
    if round_:
        b = circle(sl, x + d / 2, y + d / 2, d, fill=bg, sh=sh)
    else:
        b = rect(sl, x, y, d, d, fill=bg, r=r, sh=sh)
    icon_c(sl, name, x + d / 2, y + d / 2, d * ratio, fg)
    return b


# ---------------------------------------------------------------- composites
def chip(sl, x, y, text, bg, fg, size=10.5, h=0.3, pad=0.12, bold=True, font=FSB, line_=None, r=None):
    w = text_w(text, size, font) + 2 * pad
    rect(sl, x, y, w, h, fill=bg, line=line_, r=r if r is not None else h / 2, text=text, size=size,
         color=fg, bold=False, font=font, align="c", anchor="m", margin=0)
    return w


def chips(sl, x, y, items, bg, fg, gap=0.08, max_x=None, row_h=0.36, **kw):
    cx, cy = x, y
    for t in items:
        w = text_w(t, kw.get("size", 10.5), kw.get("font", FSB)) + 2 * kw.get("pad", 0.12)
        if max_x and cx + w > max_x and cx > x:
            cx, cy = x, cy + row_h
        chip(sl, cx, cy, t, bg, fg, **kw)
        cx += w + gap
    return cy + row_h


def background(sl, kind="light"):
    if kind == "light":
        rect(sl, 0, 0, W, H, grad=(WHITE, BG), angle=90)
    else:
        rect(sl, 0, 0, W, H, fill=kind)


def header(sl, kicker, title, sub=None, kc=SLATE, title_size=30):
    txt(sl, ML, 0.42, 9, 0.3, kicker.upper(), size=11.5, color=kc, font=FSB, spc=2.5)
    txt(sl, ML, 0.7, 11.2, 0.62, title, size=title_size, color=NAVY, bold=True, anchor="m")
    if sub:
        txt(sl, ML, 1.32, 11.5, 0.38, sub, size=14.5, color=SLATE)


def footer(sl, n, total, accent=NAVY, dark=False):
    c = "B8C3DD" if dark else MUTED
    rect(sl, ML, 7.12, 0.32, 0.035, fill=accent)
    txt(sl, ML + 0.42, 7.02, 6, 0.24, "SAPPE  ·  Daily Management System Training", size=9, color=c, anchor="m")
    txt(sl, MR - 1.2, 7.02, 1.2, 0.24, f"{n:02d} / {total}", size=9, color=c, align="r", anchor="m", font=FSB)


def notes(sl, text):
    sl.notes_slide.notes_text_frame.text = text


def card(sl, x, y, w, h, fill=WHITE, line_=LINE, r=0.14, sh=True, **kw):
    return rect(sl, x, y, w, h, fill=fill, line=line_, r=r, sh=sh, **kw)


def polar(cx, cy, r, deg):
    a = math.radians(deg)
    return cx + r * math.cos(a), cy + r * math.sin(a)
